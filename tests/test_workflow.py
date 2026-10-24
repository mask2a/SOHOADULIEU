import csv
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from sohoa.domain import HEADERS, AppError
from sohoa.application import merge, compare, JobContext
from sohoa.report import ReportWriter
from sohoa.domain import MatchResult, REVIEW, StudentRecord

def write_input(path, rows, header_row=1, headers=HEADERS):
    wb=Workbook(); ws=wb.active; ws.title='Data'
    for _ in range(header_row-1): ws.append(['Ghi chú'])
    ws.append(headers)
    if header_row==4: ws.append([f'({i})' for i in range(1,25)])
    for row in rows: ws.append(row)
    wb.save(path); wb.close()

def row(name='Nguyễn Văn An', dob='01/02/1980', year='1995', council='THCS A'):
    v=['']*24; v[0]='1'; v[4]=name; v[6]=dob; v[12]=year; v[13]=council
    return v

class WorkflowTests(unittest.TestCase):
    def test_year_named_sheet_is_read_and_lineage_preserves_name(self):
        from sohoa.excel import records, inspect
        with tempfile.TemporaryDirectory() as tmp:
            source=Path(tmp)/'THCS.CQ.1987.xlsx'
            write_input(source,[row(year='1987')],4)
            wb=load_workbook(source); wb['Data'].title='1987'; wb.save(source); wb.close()
            self.assertEqual(inspect(source,4)[:24],list(HEADERS))
            parsed=list(records(source,4,canonical=True))
            self.assertEqual(len(parsed),1)
            self.assertEqual(parsed[0].sheet,'1987')
            self.assertEqual(parsed[0].row,6)

    def test_merge_accepts_arbitrary_excel_names_and_ignores_lock_files(self):
        import json
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); data=root/'other_folder'; data.mkdir()
            write_input(data/'du lieu truong A.XLSX',[row(year='1987')],4)
            write_input(data/'ban bo sung.xlsx',[row(year='1988')],4)
            (data/'~$dang_mo.xlsx').write_text('Excel lock file',encoding='utf-8')
            (data/'notes.txt').write_text('ignore',encoding='utf-8')
            (data/'not_a_file.xlsx').mkdir()
            alias=root/'schools.csv'; alias.write_text('alias,school_id,school_name\n',encoding='utf-8')
            output=merge(data,root/'out',alias)
            wb=load_workbook(output/'1987.xlsx')
            self.assertEqual(wb['Data'].max_row,2)
            self.assertTrue((output/'1988.xlsx').exists())
            wb.close()
            self.assertEqual(json.loads((output/'pham_vi.json').read_text(encoding='utf-8'))['years_read'],[1987,1988])
            self.assertFalse((output/'bao_cao_loi.xlsx').exists())

    def test_split_folder_ids_increase_without_overwrite(self):
        from sohoa.application import run_folder
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'output'
            first=run_folder(root,'Gop_tach')
            self.assertEqual(first.name,'TACH_NAM_1')
            marker=first/'keep.txt'; marker.write_text('keep',encoding='utf-8')
            (root/'TACH_NAM_7').mkdir()
            (root/'TACH_NAM_invalid').mkdir()
            (root/'TACH_TRUONG.99').mkdir()
            second=run_folder(root,'Gop_tach')
            self.assertEqual(second.name,'TACH_NAM_8')
            self.assertEqual(marker.read_text(encoding='utf-8'),'keep')

    def test_three_schools_without_aliases_and_extra_note_on_number_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); data=root/'data'; data.mkdir()
            source=data/'THCS.CQ.1987_output.xlsx'
            write_input(source,[row(council=name,year='1987') for name in ('Phú Xuyên','Bắc Sơn','Cây Thị')],4)
            wb=load_workbook(source); wb['Data']['Y5']='Ghi chú ngoài bảng'; wb['Data']['X5']='(24'; wb.save(source); wb.close()
            alias=root/'schools.csv'; alias.write_text('alias,school_id,school_name,from_year,to_year\n',encoding='utf-8')
            output=merge(data,root/'out',alias)
            self.assertEqual({p.name for p in output.glob('*.xlsx')},{'1987.xlsx'})
            for path in output.glob('*.xlsx'):
                wb=load_workbook(path); self.assertEqual(wb['Data'].max_row,4); wb.close()

    def test_merge_compare_lineage_and_highlight(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); data=root/'data'; data.mkdir()
            write_input(data/'THCS.CQ.1995_output.xlsx',[row()],4)
            alias=root/'schools.csv'; alias.write_text('alias,school_id,school_name,from_year,to_year\nTHCS A,A,Trường A,1987,2005\n',encoding='utf-8')
            merged=merge(data,root/'out',alias)
            source=merged/'1995.xlsx'
            wb=load_workbook(source); self.assertEqual(wb.sheetnames,['Data','THONGKE']); self.assertEqual(wb['Data'].max_row,2); wb.close()
            school=root/'school.xlsx'; write_input(school,[row(dob='02/02/1980')])
            report=compare([source],school,root/'out',years={1995})
            wb=load_workbook(report/'Bao_cao_001.xlsx')
            self.assertEqual(wb.sheetnames,['Tong_hop','Lech_thong_tin','Thua_Thieu'])
            self.assertEqual(wb['Thua_Thieu']['A2'].value,REVIEW)
            self.assertEqual(wb['Lech_thong_tin']['E2'].fill.fgColor.rgb,'00FFD966')
            self.assertIn('dòng 6',wb['Thua_Thieu']['I2'].value)
            wb.close()

    def test_merge_sorts_by_school_attended_then_sequence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); data=root/'data'; data.mkdir()
            rows=[]
            for school,seq,council in [('THCS B','1','A'),('thcs a','2','Z'),('THCS A','1','Z')]:
                values=row(year='1984',council=council)
                values[11]=school; values[0]=seq; values[21]='Ghi chú gốc'; rows.append(values)
            write_input(data/'a.xlsx',rows)
            output=merge(data,root/'out')
            self.assertEqual(output.name,'TACH_NAM_1')
            wb=load_workbook(output/'1984.xlsx')
            self.assertEqual(wb.sheetnames,['Data','THONGKE'])
            self.assertEqual(list(wb['THONGKE'].values),[
                ('STT','TÊN TRƯỜNG','SỐ LƯỢNG'),
                (1,'THCS A',2),(2,'THCS B',1),('TỔNG','Tổng số trường: 2',3)])
            content=list(wb['Data'].values); exported=content[1:]; wb.close()
            self.assertEqual(content[0],(*HEADERS,'Nguồn dữ liệu'))
            self.assertEqual([(v[11],v[0]) for v in exported],[('THCS A','1'),('thcs a','2'),('THCS B','1')])
            self.assertEqual([v[24] for v in exported],[f'Tệp: a.xlsx | Sheet: Data | Dòng: {r}' for r in (4,3,2)])
            self.assertTrue(all(v[21]=='Ghi chú gốc' for v in exported))

    def test_merge_template_sheets_and_invalid_years(self):
        import sqlite3
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); data=root/'data'; data.mkdir()
            write_input(data/'a.xlsx',[row(year='1984'),row(year=''),row(year='bad')])
            wb=load_workbook(data/'a.xlsx'); wb.create_sheet('HuongDan').append(['H??ng d?n']); wb.create_sheet('DanhMuc'); wb.save(data/'a.xlsx'); wb.close()
            write_input(data/'b.xlsx',[row(year='1984'),row(year='2005')],4)
            output=merge(data,root/'out')
            wb=load_workbook(output/'1984.xlsx'); self.assertEqual(wb['Data'].max_row,3); wb.close()
            self.assertTrue((output/'2005.xlsx').exists())
            wb=load_workbook(output/'Chua_xac_dinh_nam.xlsx'); self.assertEqual(wb['Data'].max_row,3); wb.close()
            with sqlite3.connect(output/'truy_vet.sqlite') as db:
                self.assertEqual(db.execute('SELECT COUNT(*) FROM exports').fetchone()[0],5)
            db.close()
            self.assertTrue((output/'bao_cao_loi.xlsx').exists())

    def test_mapping_and_structure_rejection(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); source=root/'source.xlsx'; school=root/'school.xlsx'
            write_input(source,[row()]); write_input(school,[['Nguyễn Văn An','01/02/1980','1995']],headers=['Tên','Sinh','Năm'])
            report=compare([source],school,root/'out',mapping={4:0,6:1,12:2},years={1995})
            self.assertTrue((report/'Bao_cao_001.xlsx').exists())
            with self.assertRaises(AppError): compare([source],school,root/'out',mapping={4:0},years={1995})
            self.assertTrue(list((root/'out').glob('*/bao_cao_loi.xlsx')))

    def test_report_split_and_formula_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            v=row(name='=SUM(1,2)'); record=StudentRecord(tuple(v),'file','Data',2)
            writer=ReportWriter(tmp,'1995',3,3,row_limit=23)
            for _ in range(3): writer.add(MatchResult(REVIEW,source=record))
            paths=writer.finish(); self.assertEqual(len(paths),2)
            wb=load_workbook(paths[0]); self.assertEqual(wb['Thua_Thieu']['C2'].data_type,'s'); wb.close()

    def test_cancel_records_incomplete_state(self):
        from sohoa.domain import Cancelled
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); source=root/'source.xlsx'; write_input(source,[row()])
            with self.assertRaises(Cancelled):
                compare([source],source,root/'out',context=JobContext(cancelled=lambda:True))
            state=next((root/'out').glob('*/TRANG_THAI.txt')).read_text(encoding='utf-8')
            self.assertIn('CHƯA HOÀN TẤT',state)

if __name__=='__main__': unittest.main()
