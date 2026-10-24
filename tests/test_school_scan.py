import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from sohoa.domain import HEADERS
from sohoa.school_scan import scan_schools, suggested_group
from sohoa.application import merge

class SchoolScanTests(unittest.TestCase):
    def test_scan_review_excel_and_merge_aliases(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); data=root/'data'; data.mkdir()
            wb=Workbook(); ws=wb.active; ws.title='1987'
            for _ in range(3): ws.append([])
            ws.append(HEADERS); ws.append([f'({i})' for i in range(1,25)])
            for name,year in [('PTCS Yên Đổ','1987'),('Yên Đổ','1988'),('THCS Yên Đổ','1989'),('Yên Đổ 1','1990'),('yên đổ','1991')]:
                row=['']*24; row[4]='HỌC SINH MẪU'; row[6]='01/01/1970'; row[12]=year; row[13]=name; ws.append(row)
            wb.save(data/'sample.xlsx'); wb.close()
            output=scan_schools(data,root/'out')
            path=output/'Danh_sach_truong.xlsx'
            wb=load_workbook(path); ws=wb.active
            self.assertEqual(ws.max_row,6)
            self.assertEqual(sum(ws.cell(r,6).value for r in range(2,7)),5)
            self.assertTrue(all(ws.cell(r,2).value is None for r in range(2,7)))
            for r in range(2,7):
                if ws.cell(r,1).value!='Yên Đổ 1':
                    ws.cell(r,2,'YEN_DO'); ws.cell(r,3,'THCS Yên Đổ')
            wb.save(path); wb.close()
            merged=merge(data,root/'out',path)
            self.assertEqual({p.name for p in merged.glob('*.xlsx')},{f'{y}.xlsx' for y in range(1987,1992)})
            wb=load_workbook(merged/'1987.xlsx')
            self.assertEqual(wb['Data'].max_row,2); wb.close()

    def test_suggestion_does_not_drop_school_number(self):
        self.assertEqual(suggested_group('PTCS Yên Đổ'),suggested_group('THCS Yên Đổ'))
        self.assertNotEqual(suggested_group('Yên Đổ'),suggested_group('Yên Đổ 1'))
