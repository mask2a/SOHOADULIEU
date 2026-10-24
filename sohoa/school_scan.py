"""Inventory school labels without approving or changing school identities."""
from collections import Counter, defaultdict
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .application import JobContext, input_excels, run_folder, state
from .domain import AppError, StructureError, Cancelled, normalize
from .excel import records, save_atomic, error_report, MAX_ROWS

def suggested_group(name):
    # A hint only: preserve digits, accents and locality to avoid collapsing schools 1/2.
    return re.sub(r'^(?:TRƯỜNG\s+)?(?:THCS|PTCS|CẤP\s+II)\s+', '', normalize(name))

def scan_schools(data_dir, output, context=None):
    ctx=context or JobContext()
    files=input_excels(data_dir)
    folder=run_folder(output,'QUET_TRUONG')
    counts=Counter(); years=defaultdict(set); locations=defaultdict(set); issues=[]
    try:
        for index,path in enumerate(files,1):
            ctx.log(f'Quét tên trường: {path.name}',index,len(files))
            local=Counter(); local_years=defaultdict(set); local_locations=defaultdict(set)
            try:
                for record in records(path,4,canonical=True,log=ctx.log):
                    ctx.check()
                    name=record.values[13]  # Preserve original alias, including spacing/case.
                    local[name]+=1
                    if record.keys[2]: local_years[name].add(record.keys[2])
                    local_locations[name].add(f'{path.name} | {record.sheet}')
            except (AppError,StructureError) as exc:
                if isinstance(exc,Cancelled): raise
                issues.extend(exc.issues if isinstance(exc,StructureError) else [str(exc)])
                continue
            counts.update(local)
            for name in local:
                years[name].update(local_years[name]); locations[name].update(local_locations[name])
        if len(counts)>MAX_ROWS-1:
            raise AppError('Danh sách tên trường vượt giới hạn một sheet Excel; chia nhỏ thư mục để quét.')
        wb=Workbook(); ws=wb.active; ws.title='Danh_sach_truong'
        ws.append(['alias','school_id','school_name','from_year','to_year',
                   'so_ban_ghi','cac_nam','nguon_file_sheet','goi_y_nhom','huong_dan'])
        for name in sorted(counts,key=lambda n:(suggested_group(n),normalize(n),n)):
            ctx.check()
            values=[name,'','','','',counts[name],', '.join(sorted(years[name])),
                    '; '.join(sorted(locations[name])),suggested_group(name),
                    'Tên trống: kiểm tra dữ liệu nguồn' if not name.strip() else 'Điền cùng school_id và school_name cho các tên được xác nhận là cùng trường. Để trống cả hai nếu chưa duyệt.']
            ws.append(values)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str): cell.data_type='s'
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='305496')
        for col,width in {'A':36,'B':22,'C':38,'D':15,'E':15,'F':18,'G':35,'H':65,'I':35,'J':85}.items():
            ws.column_dimensions[col].width=width
        for row in ws.iter_rows(min_row=2,min_col=2,max_col=5):
            for cell in row: cell.fill=PatternFill('solid',fgColor='FFF2CC')
        ws.freeze_panes='B2'; ws.auto_filter.ref=ws.dimensions
        save_atomic(wb,folder/'Danh_sach_truong.xlsx')
        (folder/'HUONG_DAN.txt').write_text(
            'Duyệt Danh_sach_truong.xlsx: giữ nguyên alias; điền cùng school_id và school_name cho các tên cùng trường.\n'
            'from_year/to_year để trống sẽ áp dụng 1987–2005. Để trống cả school_id và school_name nếu chưa duyệt.\n'
            'goi_y_nhom chỉ để tham khảo; không tự gộp. Không tự bỏ số 1/2 trong tên trường.\n'
            'Lưu và đóng Excel, chọn file này ở mục Ánh xạ trường rồi chạy Gộp và tách. Không sửa dữ liệu đầu vào.\n',encoding='utf-8')
        if issues: error_report(folder/'bao_cao_loi.xlsx',issues)
        state(folder,f'{"CHƯA ĐẦY ĐỦ" if issues else "HOÀN TẤT"}: {len(counts)} tên gốc, {sum(counts.values())} bản ghi. Chưa tự gộp trường.')
        ctx.log(f'Quét xong: {len(counts)} tên gốc; xem Danh_sach_truong.xlsx trong {folder}')
        return folder
    except Exception as exc:
        state(folder,f'CHƯA HOÀN TẤT: {exc}')
        raise
