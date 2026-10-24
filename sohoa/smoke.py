"""Deterministic smoke check usable from the frozen executable."""
import json
from pathlib import Path
import tkinter as tk
from openpyxl import Workbook, load_workbook
from .domain import HEADERS
from .presentation import App
from .application import compare

def run(folder):
    folder=Path(folder).resolve(); folder.mkdir(parents=True,exist_ok=True)
    root=tk.Tk(); root.withdraw()
    app=App(root,config_dir=folder/'config')
    assert Path(app.alias.get()).is_file(), 'Missing schools CSV in frozen app'
    root.update(); root.destroy()
    record=['']*24
    record[4]='HỌC SINH MẪU'; record[6]='01/02/1980'; record[12]='1995'
    for name in ('source.xlsx','school.xlsx'):
        wb=Workbook(); ws=wb.active; ws.title='Data'; ws.append(HEADERS); ws.append(record)
        wb.save(folder/name); wb.close()
    output=compare([folder/'source.xlsx'],folder/'school.xlsx',folder/'output',years={1995})
    wb=load_workbook(output/'Bao_cao_001.xlsx')
    assert wb.sheetnames==['Tong_hop','Lech_thong_tin','Thua_Thieu']
    assert any(row[0]=='Khớp hoàn toàn' and str(row[1])=='1' for row in wb['Tong_hop'].iter_rows(values_only=True))
    wb.close()
    (folder/'result.json').write_text(json.dumps({'tkinter':'ok','excel_sqlite_report':'ok'},indent=2),encoding='utf-8')
