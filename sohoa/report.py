"""Streaming report renderer: bounded memory, three sheets per report part."""
from collections import Counter
from pathlib import Path
import pickle
import tempfile
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import PatternFill, Font
from .domain import IDENTITY, MATCH, NORMALIZED, REVIEW, text, birth
from .excel import append_text, save_atomic, MAX_ROWS

LABELS = ['Họ tên', 'Ngày sinh', 'Năm tốt nghiệp']
DETAIL_HEADERS = ['Trạng thái', 'Loại dòng', 'Họ tên chuẩn', 'Họ tên trường',
 'Ngày sinh chuẩn', 'Ngày sinh trường', 'Năm TN chuẩn', 'Năm TN trường',
 'Nguồn chuẩn', 'Nguồn trường', 'Lý do']

def detail(ws, result, source=None, candidate=False):
    source = source or result.source
    school = result.school
    values = [result.status, 'Ứng viên, chưa xác nhận' if candidate else 'Kết quả']
    marked = set()
    for k in IDENTITY:
        a = source.values[k] if source else ''
        b = school.values[k] if school else ''
        pos = len(values)
        values.extend([a,b])
        if text(a) != text(b):
            marked.update((pos,pos+1))
    reason = result.reason
    if source and school:
        notes=[]; matched=[]
        for index, (k, label) in enumerate(zip(IDENTITY, LABELS)):
            a,b = source.values[k].strip(), school.values[k].strip()
            if not a and b:
                notes.append(f'{label}: nguồn chuẩn trống; trường gửi «{b}» — đề nghị bổ sung, phải xác minh đúng học sinh trước.')
            elif a and not b:
                notes.append(f'{label}: trường thiếu thông tin; nguồn chuẩn «{a}».')
            elif not a and not b:
                notes.append(f'{label}: cả hai bên đều trống, chưa đủ căn cứ đối chiếu.')
            elif k == 6 and (not birth(a)[1] or not birth(b)[1]):
                notes.append(f'Ngày sinh chưa khớp đầy đủ/không hợp lệ: chuẩn «{a}»; trường gửi «{b}».')
            elif source.keys[index] != school.keys[index]:
                notes.append(f'Lệch {label.lower()}: chuẩn «{a}»; trường gửi «{b}».')
            else:
                matched.append(label.lower())
        if matched:
            notes.insert(0,'Khớp '+', '.join(matched)+' (sau chuẩn hóa).')
        reason=' '.join(notes + ([reason] if reason else []))
    values += [source.location if source else '', school.location if school else '', reason]
    cells = []
    for i,v in enumerate(values):
        c = WriteOnlyCell(ws, text(v)); c.data_type = 's'
        if i in marked:
            c.fill = PatternFill('solid', fgColor='FFD966')
        cells.append(c)
    ws.append(cells)

class ReportWriter:
    def __init__(self, folder, scope, source_count, school_count, row_limit=MAX_ROWS):
        self.folder = Path(folder)
        self.scope = scope
        self.source_count, self.school_count = source_count, school_count
        self.counts = Counter()
        self.part = 0
        self.limit = row_limit
        self.paths = []
        self.spool = tempfile.TemporaryFile()

    def _open(self):
        self.part += 1
        self.wb = Workbook(write_only=True)
        self.summary = self.wb.create_sheet('Tong_hop')
        self.diff = self.wb.create_sheet('Lech_thong_tin')
        self.extra = self.wb.create_sheet('Thua_Thieu')
        self.rows = 20
        for ws in (self.summary, self.diff, self.extra):
            ws.freeze_panes = 'C2'
            for col in ('A','B','C','D','E','F','G','H','I','J','K'):
                ws.column_dimensions[col].width = 28 if col not in ('I','J','K') else 55
        append_text(self.summary, ['BÁO CÁO ĐỐI SOÁT THCS', f'Phần {self.part}'])
        append_text(self.summary, ['Phạm vi năm', self.scope])
        append_text(self.summary, ['Số dòng nguồn chuẩn', self.source_count])
        append_text(self.summary, ['Số dòng trường gửi', self.school_count])
        # Pandas is restricted to small aggregate tables, never the full historical corpus.
        frame = pd.DataFrame(sorted(self.counts.items()), columns=['Trạng thái','Số lượng'])
        for status,count in frame.itertuples(index=False,name=None):
            append_text(self.summary, [status,count])
        matched = self.counts[MATCH] + self.counts[NORMALIZED]
        append_text(self.summary, ['Tỷ lệ khớp / dòng trường', f'{matched/self.school_count:.2%}' if self.school_count else 'Không có mẫu số'])
        append_text(self.summary, ['Độ bao phủ / dòng chuẩn', f'{matched/self.source_count:.2%}' if self.source_count else 'Không có mẫu số'])
        append_text(self.summary, ['Lưu ý', 'Chỉ kiểm tra họ tên, ngày sinh, năm tốt nghiệp. Khớp gồm khớp sau chuẩn hóa. Dòng ứng viên không phải số học sinh bổ sung.'])
        append_text(self.summary, [])
        append_text(self.summary, DETAIL_HEADERS)
        append_text(self.diff, DETAIL_HEADERS)
        append_text(self.extra, DETAIL_HEADERS)

    def add(self, result):
        self.counts[result.status] += 1
        pickle.dump(result, self.spool, protocol=pickle.HIGHEST_PROTOCOL)

    def _render(self, result):
        # Each candidate is a separate evidence row, not an additional student count.
        evidence = result.candidates or [result.source]
        for source in evidence:
            if self.rows >= self.limit - 1:
                self._save()
                self._open()
            detail(self.summary, result, source, bool(result.candidates))
            if result.status not in (MATCH, NORMALIZED):
                detail(self.extra, result, source, bool(result.candidates))
                if result.candidates:
                    detail(self.diff, result, source, True)
            elif result.status == NORMALIZED:
                detail(self.diff, result, source)
            self.rows += 1

    def _save(self):
        path = self.folder / f'Bao_cao_{self.part:03}.xlsx'
        save_atomic(self.wb, path)
        self.paths.append(path)

    def finish(self, check=lambda:None):
        try:
            self.spool.seek(0)
            self._open()
            while True:
                check()
                try: result = pickle.load(self.spool)
                except EOFError: break
                self._render(result)
            self._save()
            return self.paths
        finally:
            self.spool.close()
