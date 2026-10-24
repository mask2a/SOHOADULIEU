"""Two deep modules exposed as merge() and compare(); no GUI dependency."""
import csv
import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from .domain import HEADERS, IDENTITY, AppError, Cancelled, StructureError, Progress, StudentRecord, normalize
from .excel import records, append_text, save_atomic, error_report, MAX_ROWS, merge_records, add_school_statistics, setup_data_layout
from .storage import RecordStore
from .matching import reconcile
from .report import ReportWriter

@dataclass
class JobContext:
    emit: object = lambda event: None
    cancelled: object = lambda: False

    def check(self):
        if self.cancelled():
            raise Cancelled('Đã hủy; kết quả của đợt này chưa hoàn tất.')

    def log(self, message, current=0, total=0):
        self.check()
        self.emit(Progress(message, current, total))

def run_folder(output, kind):
    root = Path(output)
    root.mkdir(parents=True, exist_ok=True)
    if kind == 'Gop_tach':
        ids = [int(match.group(1)) for entry in root.iterdir()
               if (match := re.fullmatch(r'TACH_NAM_(\d+)', entry.name))]
        next_id = max(ids, default=0) + 1
        while True:
            folder = root / f'TACH_NAM_{next_id}'
            try:
                folder.mkdir()
                break
            except FileExistsError:
                next_id += 1
        (folder/'TRANG_THAI.txt').write_text('ĐANG XỬ LÝ — CHƯA HOÀN TẤT', encoding='utf-8')
        return folder
    import uuid
    folder = root / f'{kind}_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}'
    folder.mkdir()
    (folder/'TRANG_THAI.txt').write_text('ĐANG XỬ LÝ — CHƯA HOÀN TẤT', encoding='utf-8')
    return folder

def state(folder, message):
    (folder/'TRANG_THAI.txt').write_text(message, encoding='utf-8')

def aliases(path):
    """CSV columns: alias,school_id,school_name,from_year,to_year."""
    result = {}
    from contextlib import contextmanager
    from openpyxl import load_workbook

    @contextmanager
    def mapping_rows():
        if Path(path).suffix.lower() == '.xlsx':
            wb=load_workbook(path,read_only=True,data_only=True)
            try:
                if len(wb.worksheets)!=1:
                    raise AppError('Bảng ánh xạ Excel cần có duy nhất một sheet.')
                rows=wb.worksheets[0].iter_rows(values_only=True)
                headers=[str(v or '').strip() for v in next(rows,())]
                yield headers, (dict(zip(headers,row)) for row in rows)
            finally: wb.close()
        else:
            with open(path,encoding='utf-8-sig',newline='') as f:
                reader=csv.DictReader(f)
                yield reader.fieldnames or [],reader

    with mapping_rows() as (headers,reader):
        if not {'alias','school_id','school_name'}.issubset(headers):
            raise AppError('Bảng trường cần các cột alias,school_id,school_name,from_year,to_year.')
        names = {}
        for row in reader:
            sid = str(row.get('school_id') or '').strip(); name = str(row.get('school_name') or '').strip()
            alias = str(row.get('alias') or '').strip()
            if not sid and not name:
                continue  # Unreviewed rows retain direct-name grouping.
            if not sid or not name or not alias:
                raise AppError('Bảng ánh xạ có mã, tên hoặc alias trống.')
            if sid in names and names[sid] != name:
                raise AppError(f'Mã {sid} có nhiều tên trường chuẩn khác nhau.')
            names[sid] = name
            try:
                lo,hi = int(row.get('from_year') or 1987), int(row.get('to_year') or 2005)
            except ValueError as exc:
                raise AppError('Năm trong bảng ánh xạ phải là số.') from exc
            if lo > hi:
                raise AppError('Khoảng năm trong bảng ánh xạ bị đảo.')
            key = normalize(alias)
            duplicate=False
            for oldlo,oldhi,oldid,oldname in result.get(key, []):
                if max(oldlo,lo) <= min(oldhi,hi):
                    if (oldlo,oldhi,oldid,oldname)==(lo,hi,sid,name):
                        duplicate=True
                        continue
                    raise AppError(f'Alias chồng lấn năm: {row["alias"]}')
            if not duplicate:
                result.setdefault(key, []).append((lo,hi,sid,name))
    return result

def input_excels(data_dir):
    if not Path(data_dir).is_dir():
        raise AppError(f'Không tìm thấy thư mục dữ liệu: {data_dir}')
    files=sorted((p for p in Path(data_dir).iterdir()
                  if p.is_file() and p.suffix.lower() in {'.xlsx','.xlsm'}
                  and not p.name.startswith('~$') and not p.name.lower().endswith('.partial.xlsx')),
                 key=lambda p:p.name.casefold())
    if not files:
        raise AppError(f'Không tìm thấy file Excel .xlsx hoặc .xlsm trong thư mục: {data_dir}. File tạm ~$ được bỏ qua.')
    return files

def safe_name(sid, name):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', f'{sid}_{name}').strip(' .')[:100].rstrip(' .')
    if value.split('.')[0].upper() in {'CON','PRN','AUX','NUL', *[f'COM{i}' for i in range(1,10)], *[f'LPT{i}' for i in range(1,10)]}:
        value = '_' + value
    return value or 'Chua_xac_dinh'

def merge(data_dir, output, alias_file=None, context=None):
    ctx = context or JobContext()
    if not Path(data_dir).is_dir():
        raise AppError(f'Không tìm thấy thư mục dữ liệu: {data_dir}. Hãy chọn lại thư mục chứa file Excel năm.')
    folder = run_folder(output, 'Gop_tach')
    store = RecordStore(folder/'truy_vet.sqlite')
    issues=[]; unknown=0; found=set(); total_rows=0
    try:
        files = input_excels(data_dir)
        for index,path in enumerate(files,1):
            ctx.log(f'Đọc {path.name}',index,len(files))
            store.db.execute('SAVEPOINT input_file')
            file_count=0
            try:
                for record in merge_records(path,log=ctx.log):
                    ctx.check()
                    year=record.keys[2]
                    group=year if re.fullmatch(r'\d{4}',year) else '__UNKNOWN__'
                    store.add(record,'merged',group)
                    file_count+=1
                store.db.execute('RELEASE input_file')
                total_rows+=file_count
            except (StructureError,AppError) as exc:
                if isinstance(exc,Cancelled): raise
                store.db.execute('ROLLBACK TO input_file'); store.db.execute('RELEASE input_file')
                issues.extend(exc.issues if isinstance(exc,StructureError) else [str(exc)])
        found={int(year) for (year,) in store.db.execute('SELECT DISTINCT year FROM records')
               if re.fullmatch(r'\d{4}',year)}
        missing=sorted(set(range(1984,2006))-found)
        if missing:
            ctx.log('Các năm không có trong đợt này: '+', '.join(map(str,missing)))
        (folder/'pham_vi.json').write_text(json.dumps({'years_read':sorted(found),'years_absent':missing,
            'input_files':[p.name for p in files]},ensure_ascii=False,indent=2),encoding='utf-8')
        lineage=store.db
        lineage.create_function('school_sort', 1, lambda payload: normalize(json.loads(payload)[11]))
        lineage.execute('CREATE TABLE exports(output_file TEXT, output_row INTEGER, record_id INTEGER)')
        groups=[sid for (sid,) in store.db.execute('SELECT DISTINCT group_id FROM records ORDER BY group_id')]
        for sid in groups:
            ctx.log(f'Xuất năm tốt nghiệp: {sid}')
            base=sid if sid!='__UNKNOWN__' else 'Chua_xac_dinh_nam'
            wb=None; count=0; part=0
            group_count=store.db.execute('SELECT COUNT(*) FROM records WHERE group_id=?',(sid,)).fetchone()[0]
            for record_id,record in store.query('group_id=?',(sid,),order='school_sort(payload),seq,id'):
                ctx.check()
                if count==0:
                    schools={}
                    part+=1; wb=Workbook(write_only=True); ws=wb.create_sheet('Data')
                    setup_data_layout(ws)
                    append_text(ws,(*HEADERS, 'Nguồn dữ liệu'),row_number=1)
                    filename=f'{base}.xlsx' if group_count<=MAX_ROWS-1 else f'{base}_{part:03}.xlsx'
                append_text(ws,(*record.values, f'Tệp: {Path(record.file).name} | Sheet: {record.sheet} | Dòng: {record.row}'),row_number=count+2); count+=1
                school_key=normalize(record.values[11])
                school_name,school_count=schools.get(school_key,(' '.join(record.values[11].split()),0))
                schools[school_key]=(school_name,school_count+1)
                lineage.execute('INSERT INTO exports VALUES(?,?,?)',(filename,count+1,record_id))
                if sid=='__UNKNOWN__': unknown+=1
                if count==MAX_ROWS-1:
                    add_school_statistics(wb,schools)
                    save_atomic(wb,folder/filename); count=0; wb=None
            if wb is not None:
                add_school_statistics(wb,schools)
                save_atomic(wb,folder/filename)
        if unknown: issues.append(f'{unknown} dòng thiếu hoặc sai năm tốt nghiệp; xem Chua_xac_dinh_nam.xlsx.')
        if issues: error_report(folder/'bao_cao_loi.xlsx',issues)
        state(folder, f'{"CHƯA ĐẦY ĐỦ" if issues else "HOÀN TẤT"}: {total_rows} dòng; {unknown} dòng chưa xác định năm.')
        return folder
    except Exception as exc:
        state(folder,f'CHƯA HOÀN TẤT: {exc}')
        raise
    finally:
        store.close()

def _lineage(path):
    import sqlite3
    dbpath=Path(path).parent/'truy_vet.sqlite'
    if not dbpath.exists(): return {}
    db=sqlite3.connect(f'{dbpath.resolve().as_uri()}?mode=ro',uri=True)
    try:
        return {row: f'{f} | {s} | dòng {r}' for row,f,s,r in db.execute('''SELECT e.output_row,r.file,r.sheet,r.rownum FROM exports e JOIN records r ON e.record_id=r.id WHERE e.output_file=?''',(Path(path).name,))}
    finally: db.close()

def compare(source_files, school_file, output, school_header=1, mapping=None, years=None, context=None):
    ctx=context or JobContext()
    years=set(range(1984,2006)) if years is None else set(years)
    if not years or not years.issubset(set(range(1984,2006))):
        raise AppError('Phạm vi năm phải thuộc 1984–2005 và không rỗng.')
    folder=run_folder(output,'Doi_soat')
    store=RecordStore(folder/'doi_soat.sqlite')
    counts={'source':0,'school':0}; excluded={'source':0,'school':0}
    try:
        if not source_files: raise AppError('Chưa chọn file nguồn chuẩn.')
        if len({str(Path(p).resolve()).casefold() for p in source_files})!=len(source_files):
            raise AppError('Danh sách nguồn chuẩn có file trùng.')
        for side,paths,header,colmap in [('source',source_files,1,None),('school',[school_file],school_header,mapping)]:
            for path in paths:
                ctx.log(f'Đọc {side}: {Path(path).name}')
                lineage=_lineage(path) if side=='source' else {}
                for rec in records(path,header,mapping=colmap,canonical=side=='source',log=ctx.log):
                    ctx.check()
                    if years and rec.keys[2].isdigit() and int(rec.keys[2]) not in years:
                        excluded[side]+=1; continue
                    if rec.row in lineage:
                        rec=StudentRecord(rec.values,rec.file,rec.sheet,rec.row,lineage[rec.row])
                    store.add(rec,side); counts[side]+=1
            store.db.commit()
        scope=', '.join(map(str,sorted(years)))
        writer=ReportWriter(folder,scope,counts['source'],counts['school'])
        for index,result in enumerate(reconcile(store,ctx.check),1):
            writer.add(result)
            if index%500==0: ctx.log(f'Đã phân loại {index} kết quả')
        ctx.check(); writer.finish(ctx.check)
        (folder/'pham_vi.json').write_text(json.dumps({'years':scope,'counts':counts,'excluded':excluded,'columns':[HEADERS[i] for i in IDENTITY]},ensure_ascii=False,indent=2),encoding='utf-8')
        state(folder,'HOÀN TẤT — chỉ đối soát họ tên, ngày sinh, năm tốt nghiệp.')
        return folder
    except StructureError as exc:
        error_report(folder/'bao_cao_loi.xlsx',exc.issues)
        state(folder,'TỪ CHỐI ĐỐI SOÁT — yêu cầu sửa cấu trúc và gửi lại.')
        raise AppError(f'{exc}\n{folder}') from exc
    except Exception as exc:
        state(folder,f'CHƯA HOÀN TẤT: {exc}'); raise
    finally: store.close()
