"""Excel adapter: parsing, mapping validation, safe cell output."""
from datetime import date, datetime
from pathlib import Path
import re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import math
from .domain import HEADERS, IDENTITY, StudentRecord, StructureError, AppError, normalize, text

MAX_ROWS = 1_048_576

def data_sheet(wb, path):
    if wb.sheetnames == ['Data', 'THONGKE']:
        return wb['Data']
    if len(wb.worksheets) != 1:
        raise StructureError([f'{path}: cần duy nhất một sheet dữ liệu (tên bất kỳ); hiện có {wb.sheetnames}'])
    return wb.worksheets[0]


def inspect(path, header_row, sheet_name=None):
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            ws = wb[sheet_name] if sheet_name else data_sheet(wb, path)
            return [text(c.value) for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))]
        finally:
            wb.close()
    except (StructureError, AppError):
        raise
    except Exception as exc:
        raise AppError(f'Không đọc được {path}: {exc}') from exc

def auto_mapping(headers):
    available = {normalize(h): i for i, h in enumerate(headers) if h}
    return {i: available[normalize(h)] for i, h in enumerate(HEADERS) if normalize(h) in available}

def records(path, header_row, mapping=None, canonical=False, log=lambda _: None, sheet_name=None):
    headers = inspect(path, header_row, sheet_name)
    mapping = auto_mapping(headers) if mapping is None else mapping
    issues = []
    required = range(24) if canonical else IDENTITY
    for i in required:
        if i not in mapping:
            issues.append(f'{path}: thiếu ánh xạ cột {HEADERS[i]}')
    used = list(mapping.values())
    if len(used) != len(set(used)):
        issues.append(f'{path}: một cột đầu vào đang được ánh xạ cho nhiều cột chuẩn')
    for i, j in mapping.items():
        if i not in range(24) or j < 0 or j >= len(headers):
            issues.append(f'{path}: ánh xạ ngoài phạm vi {i} -> {j}')
    if canonical and tuple(normalize(h) for h in headers[:24]) != tuple(map(normalize, HEADERS)):
        issues.append(f'{path}: bộ 24 cột chuẩn không đúng thứ tự hoặc tên')
    if issues:
        raise StructureError(issues)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[sheet_name] if sheet_name else data_sheet(wb, path)
        for row_no, cells in enumerate(ws.iter_rows(min_row=header_row + 1), header_row + 1):
            raw = [c.value for c in cells]
            if not any(v is not None and text(v).strip() for v in raw):
                continue
            selected = [raw[j] for j in mapping.values() if j < len(raw)]
            nonempty = [text(v).strip() for v in selected if v is not None and text(v).strip()]
            if (canonical and header_row == 4 and row_no == 5) or (nonempty and all(re.fullmatch(r'\(\d+\)?', v) for v in nonempty)):
                log(f'Bỏ dòng số cột: {Path(path).name}:{row_no}')
                continue
            if sum(normalize(v) == normalize(HEADERS[i]) for i, j in mapping.items() if j < len(raw) for v in [raw[j]]) >= 3:
                log(f'Bỏ tiêu đề lặp: {Path(path).name}:{row_no}')
                continue
            values = [''] * 24
            for i, j in mapping.items():
                if j < len(raw):
                    v = raw[j]
                    if isinstance(v, (date, datetime)):
                        v = v.strftime('%d/%m/%Y')
                    # Recover displayed leading zeros only when Excel explicitly supplies a zero mask.
                    if i in (1, 2, 5, 22) and isinstance(v, (int, float)):
                        mask = cells[j].number_format
                        if re.fullmatch(r'0+', mask) and float(v).is_integer():
                            v = str(int(v)).zfill(len(mask))
                        else:
                            log(f'Cần kiểm tra mã dạng số, có thể mất số 0 đầu: {path}:{row_no}, {HEADERS[i]}')
                    values[i] = text(v)
            yield StudentRecord(tuple(values), str(Path(path).resolve()), ws.title, row_no)
    except (OSError, ValueError) as exc:
        raise AppError(f'Lỗi đọc {path}: {exc}') from exc
    finally:
        wb.close()


def merge_records(path, log=lambda _: None):
    """Read canonical data sheets; ignore instructions and lookup sheets."""
    wb = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    try:
        for ws in wb.worksheets:
            for row_no, cells in enumerate(ws.iter_rows(max_row=10, values_only=True), 1):
                if tuple(normalize(v) for v in cells[:24]) == tuple(map(normalize, HEADERS)):
                    sheets.append((ws.title, row_no))
                    break
    finally:
        wb.close()
    if not sheets:
        raise StructureError([f'{path}: không tìm thấy sheet có đủ 24 cột chuẩn trong 10 dòng đầu.'])
    for name, header in sheets:
        yield from records(path, header, canonical=True, log=log, sheet_name=name)

def fit_row(ws, values, row_number):
    lines = 1
    for index, value in enumerate(values, 1):
        width = ws.column_dimensions[get_column_letter(index)].width or 13
        capacity = max(1, int(width * 0.85) - 2)
        lines = max(lines, sum(max(1, math.ceil(len(part) / capacity))
                               for part in text(value).split('\n')))
    ws.row_dimensions[row_number].height = min(409, max(30, lines * 16 + 8))


def setup_data_layout(ws):
    widths = (8,20,24,30,30,22,20,12,16,18,38,42,18,38,38,22,42,30,30,24,18,42,22,38,65)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'E2'


def append_text(ws, values, row_number=None):
    from openpyxl.cell import WriteOnlyCell
    cells = []
    for value in values:
        c = WriteOnlyCell(ws, value=text(value))
        c.data_type = 's'
        if row_number is not None:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if row_number == 1:
                c.font = Font(bold=True)
                c.fill = PatternFill('solid', fgColor='D9E2F3')
        cells.append(c)
    if row_number is not None:
        fit_row(ws, values, row_number)
    ws.append(cells)

def add_school_statistics(wb, schools):
    """Summarize the records in this output part, using normalized school names."""
    from openpyxl.cell import WriteOnlyCell
    ws = wb.create_sheet('THONGKE')
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 65
    ws.column_dimensions['C'].width = 20
    ws.freeze_panes = 'A2'
    row_number = 0

    def add(values, bold=False):
        nonlocal row_number
        row_number += 1
        fit_row(ws, values, row_number)
        cells = []
        for value in values:
            cell = WriteOnlyCell(ws, value=value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.font = Font(bold=bold)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if bold:
                cell.fill = PatternFill('solid', fgColor='D9E2F3')
            cells.append(cell)
        ws.append(cells)

    add(['STT', 'TÊN TRƯỜNG', 'SỐ LƯỢNG'], bold=True)
    for index, key in enumerate(sorted(schools), 1):
        name, count = schools[key]
        add([index, name or '(Chưa có tên trường)', count])
    named_count = sum(bool(key) for key in schools)
    add(['TỔNG', f'Tổng số trường: {named_count}', sum(count for _, count in schools.values())], bold=True)


def save_atomic(wb, path):
    path = Path(path)
    tmp = path.with_name(path.stem + '.partial.xlsx')
    try:
        wb.save(tmp)
        tmp.replace(path)
    except Exception as exc:
        raise AppError(f'Không ghi được {path}: {exc}') from exc
    finally:
        wb.close()

def error_report(path, issues):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Loi_du_lieu')
    append_text(ws, ['Lỗi', 'Hướng xử lý'])
    for issue in issues:
        append_text(ws, [issue, 'Kiểm tra cấu trúc/ánh xạ; sửa và gửi lại file.'])
    save_atomic(wb, path)
